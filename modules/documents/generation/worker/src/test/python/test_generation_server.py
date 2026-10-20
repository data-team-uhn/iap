# Copyright 2026 DATA @ UHN. See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#

"""The generation worker over real HTTP, with LibreOffice stubbed unless it is installed.

A real server on a free port rather than a fake handler: reading the body is the part of this
endpoint most likely to go wrong, and only a socket shows it. The one test that converts for
real skips itself where ``soffice`` is missing, CI included.
"""

import base64
import http.client
import io
import json
import shutil
import threading
import zipfile
from http import HTTPStatus

import pytest
from docx import Document

import generation_server
import libreoffice_convert

FAKE_PDF = b"%PDF-1.7 stand-in"


def build_template(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.fixture
def server(monkeypatch):
    """A worker on a free port with two slots and no token, torn down after the test."""
    monkeypatch.delenv(generation_server.TOKEN_ENVIRONMENT_VARIABLE, raising=False)
    monkeypatch.setattr(generation_server, "_STATE", generation_server.WorkerState(2))
    instance = generation_server.DrainingHTTPServer(
        ("127.0.0.1", 0), generation_server.GenerationHandler
    )
    thread = threading.Thread(target=instance.serve_forever, daemon=True)
    thread.start()
    yield instance
    instance.shutdown()
    instance.server_close()


@pytest.fixture
def fake_soffice(monkeypatch):
    """Stand in for LibreOffice, recording what it was asked to convert."""
    converted = []

    def docx_to_pdf(source, output_dir, log=None):
        converted.append(source.read_bytes())
        target = output_dir / f"{source.stem}.pdf"
        target.write_bytes(FAKE_PDF)
        return target

    monkeypatch.setattr(libreoffice_convert, "docx_to_pdf", docx_to_pdf)
    return converted


def post(server, body: bytes, headers: dict | None = None, path: str = "/generate"):
    """POST ``body`` and return the status and the decoded JSON reply."""
    connection = http.client.HTTPConnection(*server.server_address, timeout=30)
    try:
        connection.request("POST", path, body=body, headers=headers or {})
        response = connection.getresponse()
        return response.status, json.loads(response.read() or b"{}")
    finally:
        connection.close()


def generate_request(template: bytes, context: dict) -> bytes:
    return json.dumps(
        {"template": base64.b64encode(template).decode("ascii"), "context": context}
    ).encode()


class TestGenerate:
    def test_answers_with_the_filled_docx_and_its_pdf(self, server, fake_soffice):
        status, reply = post(
            server, generate_request(build_template("Dear {{ name }},"), {"name": "Dr. A"})
        )

        assert status == HTTPStatus.OK
        assert base64.b64decode(reply["pdf"]) == FAKE_PDF
        docx = Document(io.BytesIO(base64.b64decode(reply["docx"])))
        assert docx.paragraphs[0].text == "Dear Dr. A,"
        assert fake_soffice == [base64.b64decode(reply["docx"])], (
            "the PDF must be made from the docx that is sent back"
        )

    def test_a_template_mistake_is_a_422_that_says_what_is_wrong(self, server, fake_soffice):
        status, reply = post(server, generate_request(build_template("{{ missing }}"), {}))

        assert status == HTTPStatus.UNPROCESSABLE_ENTITY
        assert "missing" in reply["error"]
        assert fake_soffice == [], "nothing is converted for a template that did not render"

    def test_a_zip_bomb_is_refused_before_it_is_opened(self, server, fake_soffice, monkeypatch):
        monkeypatch.setenv("IAP_MAX_EXPANDED_BYTES", "1000")
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bomb:
            bomb.writestr("word/document.xml", b"0" * 100_000)

        status, reply = post(server, generate_request(archive.getvalue(), {}))

        assert status == HTTPStatus.UNPROCESSABLE_ENTITY
        assert "expands" in reply["error"]

    def test_a_conversion_failure_keeps_its_details_in_the_log(self, server, monkeypatch):
        def failing(source, output_dir, log=None):
            raise RuntimeError("soffice said something about /tmp/iap-gen-secret")

        monkeypatch.setattr(libreoffice_convert, "docx_to_pdf", failing)
        status, reply = post(server, generate_request(build_template("x"), {}))

        assert status == HTTPStatus.INTERNAL_SERVER_ERROR
        assert reply["reference"]
        assert "iap-gen-secret" not in json.dumps(reply), "internals must stay in the log"

    def test_is_busy_when_every_slot_is_taken(self, server, fake_soffice):
        state = generation_server._STATE
        for _ in range(state.slot_count):
            state.slots.acquire()
        try:
            status, reply = post(server, generate_request(build_template("x"), {}))
        finally:
            for _ in range(state.slot_count):
                state.slots.release()

        assert status == HTTPStatus.SERVICE_UNAVAILABLE
        assert "busy" in reply["error"]
        assert fake_soffice == []


class TestBadRequests:
    @pytest.mark.parametrize(
        "body, fragment",
        [
            (b"not json", "JSON object"),
            (b"[]", "JSON object"),
            (b'{"context": {}}', "template"),
            (b'{"template": "%%%", "context": {}}', "base64"),
            (b'{"template": "eA==", "context": []}', "context"),
        ],
    )
    def test_a_malformed_body_is_a_400(self, server, fake_soffice, body, fragment):
        status, reply = post(server, body)
        assert status == HTTPStatus.BAD_REQUEST
        assert fragment in reply["error"]

    def test_a_body_over_the_limit_is_refused_unread(self, server, fake_soffice, monkeypatch):
        monkeypatch.setenv(generation_server.MAX_REQUEST_BYTES_VARIABLE, "100")
        status, _reply = post(server, generate_request(build_template("x"), {}))
        assert status == HTTPStatus.REQUEST_ENTITY_TOO_LARGE

    def test_an_unknown_path_is_a_404(self, server):
        status, _reply = post(server, b"", path="/parse")
        assert status == HTTPStatus.NOT_FOUND


class TestGuards:
    def test_a_browser_origin_is_refused(self, server, fake_soffice):
        status, _reply = post(
            server,
            generate_request(build_template("x"), {}),
            headers={"Origin": "https://evil.example"},
        )
        assert status == HTTPStatus.FORBIDDEN
        assert fake_soffice == []

    def test_the_token_is_required_once_set(self, server, fake_soffice, monkeypatch):
        monkeypatch.setenv(generation_server.TOKEN_ENVIRONMENT_VARIABLE, "s3cret")
        body = generate_request(build_template("x"), {})

        assert post(server, body)[0] == HTTPStatus.UNAUTHORIZED
        assert post(server, body, headers={"Authorization": "Bearer s3cret"})[0] == HTTPStatus.OK

    def test_the_docling_token_does_not_open_this_worker(self, server, fake_soffice, monkeypatch):
        monkeypatch.setenv(generation_server.TOKEN_ENVIRONMENT_VARIABLE, "generation-secret")
        monkeypatch.setenv("IAP_DOCLING_TOKEN", "docling-secret")
        status, _reply = post(
            server,
            generate_request(build_template("x"), {}),
            headers={"Authorization": "Bearer docling-secret"},
        )
        assert status == HTTPStatus.UNAUTHORIZED

    def test_health_needs_no_token(self, server, monkeypatch):
        monkeypatch.setenv(generation_server.TOKEN_ENVIRONMENT_VARIABLE, "s3cret")
        connection = http.client.HTTPConnection(*server.server_address, timeout=30)
        try:
            connection.request("GET", "/health")
            response = connection.getresponse()
            reply = json.loads(response.read())
        finally:
            connection.close()
        assert response.status == HTTPStatus.OK
        assert reply == {"status": "ok", "ready": True, "slots": 2}


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
class TestRealConversion:
    def test_produces_a_pdf(self, server):
        status, reply = post(
            server, generate_request(build_template("Dear {{ name }},"), {"name": "Dr. A"})
        )
        assert status == HTTPStatus.OK
        assert base64.b64decode(reply["pdf"]).startswith(b"%PDF")
