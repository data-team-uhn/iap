# Copyright 2026 DATA @ UHN. See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#

"""Filling a template: values in, formatting kept, and every template mistake a TemplateError.

The templates are built here with python-docx rather than kept as binary fixtures, so what a
test feeds the renderer is readable next to what it expects back.
"""

import io

import pytest
from docx import Document

import render_template


def build_template(*paragraphs: str, table: list[list[str]] | None = None, bold: str = "") -> bytes:
    """A .docx with one paragraph per string, an optional table, and an optional bold run."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if bold:
        document.add_paragraph().add_run(bold).bold = True
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def text_of(rendered: bytes) -> list[str]:
    """The paragraphs and table cells of a rendered .docx, in order."""
    document = Document(io.BytesIO(rendered))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return lines


class TestRendering:
    def test_fills_the_placeholders(self):
        rendered = render_template.render(
            build_template("Dear {{ submitter.name }},"), {"submitter": {"name": "Dr. A"}}
        )
        assert "Dear Dr. A," in text_of(rendered)

    def test_escapes_values_so_the_document_still_opens(self):
        # Without autoescape this writes a raw & into the XML and Word refuses the file
        rendered = render_template.render(
            build_template("{{ title }}"), {"title": "A & B <combined>"}
        )
        assert "A & B <combined>" in text_of(rendered)

    def test_keeps_the_formatting_of_the_tag(self):
        rendered = render_template.render(build_template(bold="{{ name }}"), {"name": "Dr. A"})
        runs = [run for p in Document(io.BytesIO(rendered)).paragraphs for run in p.runs]
        assert any(run.text == "Dr. A" and run.bold for run in runs), (
            "a bold tag must produce a bold value"
        )

    def test_repeats_a_table_row_per_item(self):
        template = build_template(
            table=[
                ["Reviewer", "Role"],
                ["{%tr for r in reviewers %}", ""],
                ["{{ r.name }}", "{{ r.role }}"],
                ["{%tr endfor %}", ""],
            ]
        )
        reviewers = [{"name": "X", "role": "Chair"}, {"name": "Y", "role": "Member"}]
        rendered = render_template.render(template, {"reviewers": reviewers})
        assert text_of(rendered) == ["Reviewer", "Role", "X", "Chair", "Y", "Member"]

    def test_drops_a_conditional_paragraph(self):
        template = build_template("{%p if conditions %}", "Subject to conditions.", "{%p endif %}")
        rendered = render_template.render(template, {"conditions": False})
        assert "Subject to conditions." not in text_of(rendered)


class TestTemplateMistakes:
    """Every way a template can be wrong is reported as one, for the 422."""

    def test_a_missing_value_names_it(self):
        with pytest.raises(render_template.TemplateError, match="submitter"):
            render_template.render(build_template("Dear {{ submitter.name }},"), {})

    def test_a_row_tag_outside_a_table_is_a_template_error(self):
        # The mistake the first real template made: a table drawn with | characters
        with pytest.raises(render_template.TemplateError, match="tr"):
            render_template.render(build_template("| {%tr for r in reviewers %} |"), {})

    def test_a_template_cannot_reach_python_internals(self):
        with pytest.raises(render_template.TemplateError):
            render_template.render(
                build_template("{{ name.__class__.__mro__ }}"), {"name": "Dr. A"}
            )

    def test_bytes_that_are_not_a_docx_are_a_template_error(self):
        with pytest.raises(render_template.TemplateError, match="not a readable .docx"):
            render_template.render(b"not a zip at all", {})
